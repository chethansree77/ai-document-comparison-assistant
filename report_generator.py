
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO
from datetime import datetime


def generate_comparison_report(
    document_a_name,
    document_b_name,
    comparison_result,
):
    """
    Generate a downloadable Word document containing
    the AI comparison of two PDF documents.
    """

    doc = Document()

    # Configure page margins
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # Configure default font
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    # Report title
    title = doc.add_heading(
        "AI Document Comparison Report", level=0
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Report information
    doc.add_paragraph(
        f"Generated on: {datetime.now().strftime('%d-%m-%Y %H:%M')}"
    )
    doc.add_paragraph(f"Document A: {document_a_name}")
    doc.add_paragraph(f"Document B: {document_b_name}")

    doc.add_heading("Comparison Results", level=1)

    # Add comparison text while preserving headings
    # and basic Markdown formatting.
    for line in comparison_result.splitlines():
        stripped_line = line.strip()

        if not stripped_line:
            continue

        if stripped_line.startswith("### "):
            doc.add_heading(stripped_line[4:], level=3)

        elif stripped_line.startswith("## "):
            doc.add_heading(stripped_line[3:], level=2)

        elif stripped_line.startswith("# "):
            doc.add_heading(stripped_line[2:], level=1)

        elif stripped_line.startswith("- ") or stripped_line.startswith("* "):
            doc.add_paragraph(
                stripped_line[2:],
                style="List Bullet",
            )

        elif len(stripped_line) > 2 and stripped_line[0].isdigit() and stripped_line[1:3] == ". ":
            doc.add_paragraph(
                stripped_line[3:],
                style="List Number",
            )

        else:
            doc.add_paragraph(stripped_line)

    # Save the Word document in memory
    output = BytesIO()
    doc.save(output)
    output.seek(0)

    return output.getvalue()
